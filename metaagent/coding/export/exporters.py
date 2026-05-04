from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def export_records(
    records: Iterable[dict],
    out_dir: Path,
    basename: str = "coding_sheet",
    index_records: Optional[Iterable[dict]] = None,
    column_order: Optional[list[str]] = None,
    index_column_order: Optional[list[str]] = None,
) -> tuple[Path, Optional[Path]]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_path = out_dir / f"{basename}_{ts}.xlsx"
    records_list = list(records)
    index_list = list(index_records) if index_records else []
    if not records_list and not index_list:
        pd.DataFrame().to_excel(xlsx_path, index=False)
        return xlsx_path, None

    df = pd.DataFrame(records_list) if records_list else pd.DataFrame()
    df_index = pd.DataFrame(index_list) if index_list else pd.DataFrame()

    if column_order and not df.empty:
        cols = [c for c in column_order if c in df.columns]
        rest = [c for c in df.columns if c not in cols]
        df = df[cols + rest]

    if index_column_order and not df_index.empty:
        cols = [c for c in index_column_order if c in df_index.columns]
        rest = [c for c in df_index.columns if c not in cols]
        df_index = df_index[cols + rest]

    # Write Excel with minimal formatting
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as w:
        if not df.empty:
            df.to_excel(w, index=False, sheet_name="All Records")
            _format_excel_sheet(w.book["All Records"], df)
        if not df_index.empty:
            df_index.to_excel(w, index=False, sheet_name="Index")
            _format_excel_sheet(w.book["Index"], df_index)

    return xlsx_path, None


def _format_excel_sheet(ws, df):
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    ALT_ROW_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    BORDER = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )

    # Header
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    # Alternating row colors
    for row_idx in range(2, len(df) + 2):
        if row_idx % 2 == 0:
            for cell in ws[row_idx]:
                cell.fill = ALT_ROW_FILL
                cell.border = BORDER
        else:
            for cell in ws[row_idx]:
                cell.border = BORDER
