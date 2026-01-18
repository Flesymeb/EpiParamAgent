#!/usr/bin/env python3
"""
比较Excel筛选结果与Ground Truth，计算混淆矩阵和性能指标

使用方法:
    python compare_with_ground_truth.py --excel test/search_v2_raw_results.xlsx --ground-truth ../langgraph_runs/ground_truth/search_v2/search_v2_gt.csv --pool ../langgraph_runs/ground_truth/search_v2/search_v2_raw.csv
"""

import argparse
import csv
import sys
from pathlib import Path
from difflib import SequenceMatcher

try:
    import pandas as pd

    USE_PANDAS = True
except ImportError:
    try:
        import openpyxl

        USE_PANDAS = False
    except ImportError:
        print("❌ 错误: 需要安装 pandas 或 openpyxl")
        print("   请运行: pip install pandas openpyxl")
        sys.exit(1)


def normalize_title(title: str) -> str:
    """标准化标题用于匹配"""
    if not title:
        return ""
    # 转小写，移除标点和多余空格
    import re

    title = title.lower()
    title = re.sub(r"[^\w\s]", " ", title)
    title = " ".join(title.split())
    return title


def title_similarity(title1: str, title2: str) -> float:
    """计算两个标题的相似度"""
    norm1 = normalize_title(title1)
    norm2 = normalize_title(title2)
    return SequenceMatcher(None, norm1, norm2).ratio()


def load_ground_truth(gt_file: Path) -> dict:
    """加载Ground Truth，返回 {normalized_title: original_title}"""
    gt_titles = {}

    with open(gt_file, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            # 尝试多种可能的列名
            title = None
            for key in ["Title", "title", "TITLE"]:
                if key in row:
                    title = row[key].strip()
                    break

            if title:
                norm_title = normalize_title(title)
                gt_titles[norm_title] = title

    print(f"📋 Ground Truth: {len(gt_titles)} 篇文献")
    return gt_titles


def load_pool(pool_file: Path) -> dict:
    """加载筛选池，返回 {normalized_title: original_title}"""
    pool_titles = {}

    with open(pool_file, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            # 尝试多种可能的列名
            title = None
            for key in ["Title", "title", "TITLE"]:
                if key in row:
                    title = row[key].strip()
                    break

            if title:
                norm_title = normalize_title(title)
                pool_titles[norm_title] = title

    print(f"📋 筛选池: {len(pool_titles)} 篇文献")
    return pool_titles


def load_excel_results(excel_file: Path) -> dict:
    """
    加载Excel筛选结果
    返回 {normalized_title: decision}

    假设Excel格式:
    - 第一列: Title
    - 某列包含筛选决策 (Include/Exclude 或类似标记)
    """

    if USE_PANDAS:
        # 使用pandas读取
        df = pd.read_excel(excel_file)
        headers = df.columns.tolist()
    else:
        # 使用openpyxl读取
        wb = openpyxl.load_workbook(excel_file)
        ws = wb.active

        # 读取表头
        headers = []
        for cell in ws[1]:
            headers.append(cell.value)

    print(f"\n📊 Excel 表头: {headers}")

    # 寻找关键列
    title_col = None
    decision_col = None

    for i, header in enumerate(headers):
        if header:
            header_lower = str(header).lower()
            if "title" in header_lower:
                title_col = i
            elif any(
                keyword in header_lower
                for keyword in [
                    "decision",
                    "include",
                    "exclude",
                    "relevant",
                    "select",
                    "filter",
                    "screening",
                ]
            ):
                decision_col = i

    if title_col is None:
        print(f"❌ 错误: 未找到Title列")
        print(f"   可用列: {headers}")
        sys.exit(1)

    if decision_col is None:
        print(f"⚠️  警告: 未找到筛选决策列，请手动指定")
        print(f"   可用列: {headers}")
        # 让用户选择
        print(f"\n请输入筛选决策列的索引 (0-{len(headers)-1}):")
        for i, h in enumerate(headers):
            print(f"  [{i}] {h}")
        try:
            decision_col = int(input("列索引: "))
        except:
            print("❌ 无效输入")
            sys.exit(1)

    print(f"✅ Title列: {headers[title_col]} (索引 {title_col})")
    print(f"✅ 决策列: {headers[decision_col]} (索引 {decision_col})")

    # 读取数据
    results = {}
    included_count = 0
    excluded_count = 0

    if USE_PANDAS:
        # 使用pandas处理
        title_col_name = headers[title_col]
        decision_col_name = headers[decision_col]

        for idx, row in df.iterrows():
            title = row[title_col_name]
            decision = row[decision_col_name]

            if pd.isna(title) or not title:
                continue

            norm_title = normalize_title(str(title))

            # 判断是Include还是Exclude
            decision_str = str(decision).lower() if not pd.isna(decision) else ""

            # 多种可能的表达方式
            is_included = False
            if any(
                keyword in decision_str
                for keyword in [
                    "include",
                    "yes",
                    "relevant",
                    "accept",
                    "true",
                    "1",
                    "select",
                ]
            ):
                is_included = True
            elif any(
                keyword in decision_str
                for keyword in ["exclude", "no", "irrelevant", "reject", "false", "0"]
            ):
                is_included = False
            else:
                # 如果无法判断，打印出来让用户确认
                if idx <= 5:  # 只打印前几行作为示例
                    print(
                        f"   行 {idx+2}: '{str(title)[:50]}...' => 决策: '{decision}'"
                    )

            results[norm_title] = is_included

            if is_included:
                included_count += 1
            else:
                excluded_count += 1
    else:
        # 使用openpyxl处理
        for row_idx, row in enumerate(
            ws.iter_rows(min_row=2, values_only=True), start=2
        ):
            if not row or len(row) <= max(title_col, decision_col):
                continue

            title = row[title_col]
            decision = row[decision_col]

            if not title:
                continue

            norm_title = normalize_title(str(title))

            # 判断是Include还是Exclude
            decision_str = str(decision).lower() if decision else ""

            # 多种可能的表达方式
            is_included = False
            if any(
                keyword in decision_str
                for keyword in [
                    "include",
                    "yes",
                    "relevant",
                    "accept",
                    "true",
                    "1",
                    "select",
                ]
            ):
                is_included = True
            elif any(
                keyword in decision_str
                for keyword in ["exclude", "no", "irrelevant", "reject", "false", "0"]
            ):
                is_included = False
            else:
                # 如果无法判断，打印出来让用户确认
                if row_idx <= 5:  # 只打印前几行作为示例
                    print(f"   行 {row_idx}: '{title[:50]}...' => 决策: '{decision}'")

            results[norm_title] = is_included

            if is_included:
                included_count += 1
            else:
                excluded_count += 1

    print(f"\n📋 Excel筛选结果:")
    print(f"   ✅ Include: {included_count} 篇")
    print(f"   ❌ Exclude: {excluded_count} 篇")
    print(f"   📊 总计: {len(results)} 篇")

    return results


def match_titles(
    excel_results: dict, gt_titles: set, pool_titles: set, threshold: float = 0.9
):
    """
    匹配Excel结果与Ground Truth

    返回: (TP, FP, FN, TN, matched_results, unmatched)
    """
    # 转换为标准化标题集合
    gt_norm = {normalize_title(t) for t in gt_titles}
    pool_norm = {normalize_title(t) for t in pool_titles}

    TP = 0  # Excel=Include, GT=Include
    FP = 0  # Excel=Include, GT=Exclude
    FN = 0  # Excel=Exclude, GT=Include
    TN = 0  # Excel=Exclude, GT=Exclude

    matched = []
    unmatched = []

    for norm_title, is_included in excel_results.items():
        # 检查是否在GT中
        in_gt = norm_title in gt_norm

        # 检查是否在筛选池中
        if norm_title not in pool_norm:
            # 模糊匹配
            best_match = None
            best_score = 0

            for pool_title in pool_norm:
                score = SequenceMatcher(None, norm_title, pool_title).ratio()
                if score > best_score:
                    best_score = score
                    best_match = pool_title

            if best_score >= threshold:
                norm_title = best_match
                in_gt = best_match in gt_norm
            else:
                unmatched.append((norm_title, is_included, best_score))
                continue

        # 统计混淆矩阵
        if is_included and in_gt:
            TP += 1
            matched.append((norm_title, "TP"))
        elif is_included and not in_gt:
            FP += 1
            matched.append((norm_title, "FP"))
        elif not is_included and in_gt:
            FN += 1
            matched.append((norm_title, "FN"))
        else:  # not is_included and not in_gt
            TN += 1
            matched.append((norm_title, "TN"))

    return TP, FP, FN, TN, matched, unmatched


def print_confusion_matrix(TP, FP, FN, TN):
    """打印混淆矩阵和性能指标"""
    print("\n" + "=" * 70)
    print("📊 混淆矩阵分析")
    print("=" * 70)

    print("\n混淆矩阵:")
    print("┌─────────────────────────┬──────────────────┬──────────────────┐")
    print("│                         │   实际相关(GT)   │   实际不相关     │")
    print("├─────────────────────────┼──────────────────┼──────────────────┤")
    print(f"│ Excel判定为相关         │   TP = {TP:<8} │   FP = {FP:<8} │")
    print("├─────────────────────────┼──────────────────┼──────────────────┤")
    print(f"│ Excel判定为不相关       │   FN = {FN:<8} │   TN = {TN:<8} │")
    print("└─────────────────────────┴──────────────────┴──────────────────┘")

    # 计算性能指标
    total = TP + FP + FN + TN
    sensitivity = TP / (TP + FN) if (TP + FN) > 0 else 0
    specificity = TN / (TN + FP) if (TN + FP) > 0 else 0
    precision = TP / (TP + FP) if (TP + FP) > 0 else 0
    npv = TN / (TN + FN) if (TN + FN) > 0 else 0
    f1 = (
        2 * precision * sensitivity / (precision + sensitivity)
        if (precision + sensitivity) > 0
        else 0
    )
    accuracy = (TP + TN) / total if total > 0 else 0

    print("\n📊 性能指标:")
    print("┌─────────────────────────────┬─────────┬─────────────────────────┐")
    print("│ 指标                        │  数值   │  计算公式               │")
    print("├─────────────────────────────┼─────────┼─────────────────────────┤")
    print(
        f"│ 📈 Sensitivity (召回率)     │ {sensitivity:6.1%} │ TP/(TP+FN) = {TP}/{TP+FN:<8} │"
    )
    print(
        f"│ 🎯 Specificity (特异度)     │ {specificity:6.1%} │ TN/(TN+FP) = {TN}/{TN+FP:<8} │"
    )
    print(
        f"│ 💎 Precision (精确率)       │ {precision:6.1%} │ TP/(TP+FP) = {TP}/{TP+FP:<8} │"
    )
    print(
        f"│ ✅ NPV (阴性预测值)         │ {npv:6.1%} │ TN/(TN+FN) = {TN}/{TN+FN:<8} │"
    )
    print(f"│ 🏆 F1-score                 │ {f1:6.1%} │ 2×P×R/(P+R)             │")
    print(
        f"│ ✨ Accuracy (准确率)        │ {accuracy:6.1%} │ (TP+TN)/Total = {TP+TN}/{total:<4} │"
    )
    print("└─────────────────────────────┴─────────┴─────────────────────────┘")


def main():
    parser = argparse.ArgumentParser(description="比较Excel筛选结果与Ground Truth")
    parser.add_argument(
        "--excel",
        type=Path,
        required=True,
        help="Excel筛选结果文件",
    )
    parser.add_argument(
        "--ground-truth",
        type=Path,
        required=True,
        help="Ground Truth CSV文件",
    )
    parser.add_argument(
        "--pool",
        type=Path,
        required=True,
        help="原始筛选池CSV文件",
    )
    parser.add_argument(
        "--threshold",
        type=float,
        default=0.9,
        help="标题匹配相似度阈值 (默认0.9)",
    )

    args = parser.parse_args()

    # 检查文件存在
    if not args.excel.exists():
        print(f"❌ 错误: Excel文件不存在: {args.excel}")
        sys.exit(1)

    if not args.ground_truth.exists():
        print(f"❌ 错误: Ground Truth文件不存在: {args.ground_truth}")
        sys.exit(1)

    if not args.pool.exists():
        print(f"❌ 错误: 筛选池文件不存在: {args.pool}")
        sys.exit(1)

    print("=" * 70)
    print("📊 Excel筛选结果 vs Ground Truth 对比分析")
    print("=" * 70)

    # 加载数据
    gt_titles = load_ground_truth(args.ground_truth)
    pool_titles = load_pool(args.pool)
    excel_results = load_excel_results(args.excel)

    # 匹配和分析
    TP, FP, FN, TN, matched, unmatched = match_titles(
        excel_results,
        set(gt_titles.values()),
        set(pool_titles.values()),
        args.threshold,
    )

    # 打印结果
    print_confusion_matrix(TP, FP, FN, TN)

    # 打印未匹配的记录
    if unmatched:
        print(f"\n⚠️  警告: {len(unmatched)} 条记录无法匹配到筛选池:")
        for i, (title, decision, score) in enumerate(unmatched[:5], 1):
            print(f"   {i}. {title[:60]}... (最佳匹配度: {score:.2f})")
        if len(unmatched) > 5:
            print(f"   ... 还有 {len(unmatched) - 5} 条")


if __name__ == "__main__":
    main()
